#!/usr/bin/env python3
from __future__ import annotations

import json
import math
import re
import subprocess
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any

from docx import Document


ROOT = Path("/Users/wanghaoyu/Desktop/当前工作区/AIGC影视/structured_cases/chi_qingming")
SCRIPT_PATH = Path("/Users/wanghaoyu/Desktop/AI短片《池清明到底是谁》 剧本-终稿.docx")
VIDEO_PATH = Path("/Users/wanghaoyu/Desktop/杰克十七号_AI短片电影_《池清明到底是谁》.mp4")
DATASET_PATH = ROOT / "chi_qingming.dataset.json"
KEYFRAME_DIR = ROOT / "keyframes"
SCENE_THRESHOLD = 0.40
MIN_SHOT_SECONDS = 0.5
LONG_SHOT_REVIEW_SECONDS = 12.0


SCENE_HEADING_RE = re.compile(
    r"^(?P<number>\d+[A-Z]?)\s+(?P<location>.+?)\s+(?P<time_of_day>日|夜|昏|晨|早|晚|凌晨|傍晚)\s+(?P<interior_exterior>内|外)$"
)


CHARACTER_TERMS = {
    "wu_yijie": {
        "name": "吴翊杰",
        "aliases": ["吴翊杰", "小杰"],
        "description": "初三学生，池清明的朋友，执着追查池清明消失与世界规则异常，最终选择冲向天台外的透明薄膜。",
        "prompt": "Chinese middle school boy, black-white school uniform, anxious but determined expression, cinematic suspense short film style",
    },
    "chi_qingming": {
        "name": "池清明",
        "aliases": ["池清明", "大池子"],
        "description": "吴翊杰的朋友，被集体记忆抹除的人；通过篮球隐喻提出“忘掉篮筐”以跳出规则。",
        "prompt": "mysterious Chinese middle school boy, basketball court memory, confident and detached, cinematic backlight",
    },
    "shi_qi": {
        "name": "石琦",
        "aliases": ["石琦"],
        "description": "班级同学，起初冷漠旁观，实际记得被抹除者并掌握部分规则，最终通过广播帮助吴翊杰。",
        "prompt": "quiet Chinese middle school girl, school uniform with sleeve covers, restrained expression, secretive observer",
    },
    "teacher_yan": {
        "name": "闫老师",
        "aliases": ["闫老师", "班主任闫老师"],
        "description": "班主任和规则维护者形象，反复要求吴翊杰回到考试、升学和服从的轨道。",
        "prompt": "stern Chinese homeroom teacher, school office or rooftop, authoritative posture, unsettling calm",
    },
    "daguang": {
        "name": "大光",
        "aliases": ["大光", "寸头体育生"],
        "description": "体育生，教室里占据原本属于池清明的座位，后来在校园广场阻拦吴翊杰去天台。",
        "prompt": "buzz-cut athletic Chinese student, red basketball jersey, holding basketball, impatient and forceful",
    },
    "twin_tail_girl": {
        "name": "双马尾女生",
        "aliases": ["双马尾女生"],
        "description": "走廊里被吴翊杰追问的同学，后在主楼门口被撞倒并试图拖住吴翊杰。",
        "prompt": "Chinese schoolgirl with twin ponytails, carrying books, hurt ankle, accusatory expression",
    },
    "glasses_boy": {
        "name": "眼镜男",
        "aliases": ["眼镜男"],
        "description": "走廊里被吴翊杰抓住追问的同学，惊恐地看着他。",
        "prompt": "Chinese male student with glasses, frightened reaction in crowded school corridor",
    },
    "bun_hair_girl": {
        "name": "丸子头女生",
        "aliases": ["丸子头女生"],
        "description": "走廊里护住双马尾女生的同学，代表旁观者对吴翊杰的排斥。",
        "prompt": "Chinese schoolgirl with bun hairstyle, protective posture, crowded school corridor",
    },
    "middle_aged_woman_voice": {
        "name": "中年女声",
        "aliases": ["中年女声", "刘莉", "刘阿姨"],
        "description": "电话亭段落中接听池清明家中座机的女性声音，否认认识池清明且称自己没有孩子。",
        "prompt": "off-screen middle aged woman voice on telephone, ambiguous memory-erasure clue",
    },
}


SCENE_ASSETS = {
    "school_office": ("校办公室", "日间教师办公室，堆满试卷和讲义，茶杯、抽屉、办公桌构成考试压力与管控气氛。"),
    "internet_cafe": ("网吧", "夜间网吧，电脑屏幕显示 QQ、blog 和 404 页面，是池清明网络身份被抹除的证据空间。"),
    "phone_booth": ("电话亭", "夜间电话亭，吴翊杰用通讯本拨打池清明家电话，城市车辆从前景掠过。"),
    "basketball_court": ("学校篮球场", "日间校园篮球场，池清明用闭眼投篮讲述跳出规则的方法，是核心隐喻场景。"),
    "corridor_classroom": ("学校走廊/教室", "日间教室和走廊，池清明的座位被大光替代，周围同学集体否认他的存在。"),
    "library": ("校图书馆", "惨白光柱和尘埃中的图书馆，借阅记录揭示集体无意识的记忆错位。"),
    "campus_square": ("校园广场", "昏色校园广场，吴翊杰奔向天台时被大光以训练名义阻拦。"),
    "main_hall": ("校园主楼·门口大厅", "昏色主楼入口大厅，双马尾女生摔倒，形成道德与规则的阻拦。"),
    "stairwell": ("校园主楼·楼梯间", "被旧课桌椅和人群堵塞的楼梯间，石琦广播制造缺口。"),
    "rooftop": ("教学楼天台", "昏色教学楼天台，废旧桌椅、蜻蜓、透明薄膜和闫老师构成最终冲突空间。"),
}


PROP_TERMS = {
    "phone": ("手机", ["手机"], "被闫老师没收的手机，也是吴翊杰联系池清明失败的线索。"),
    "exam_papers": ("试卷/讲义", ["试卷", "讲义", "卷子", "空白卷子"], "办公室中成堆的学习材料，象征考试规训。"),
    "tea_cup": ("茶杯", ["茶杯", "杯子", "茶水"], "闫老师训斥时碰倒，弄湿吴翊杰的卷子。"),
    "contact_book": ("通讯本", ["通讯本"], "吴翊杰在手机被没收后用来拨打池清明家电话。"),
    "basketball": ("篮球", ["篮球", "投篮", "篮筐"], "池清明讲述世界规则的核心隐喻道具。"),
    "qq_blog_404": ("QQ/blog/404 页面", ["QQ", "blog", "404"], "网吧中显示池清明线上身份被注销或查无此人的证据。"),
    "borrowing_record": ("借阅记录纸", ["借阅记录纸", "借阅登记"], "图书馆中记录被抹除学生姓名的证据。"),
    "library_rules_book": ("《借阅守则》", ["借阅守则"], "石琦藏匿借阅记录纸的厚重书籍。"),
    "school_broadcast": ("广播", ["广播", "麦克风", "音响"], "石琦通过全校广播念出被抹除者姓名，打断楼梯间阻拦。"),
    "old_desks_chairs": ("旧课桌椅", ["旧课桌椅", "课桌", "桌椅"], "楼梯间和天台上阻拦或构成空间障碍的物件。"),
    "transparent_membrane": ("透明薄膜", ["透明薄膜", "气泡", "玻璃污渍"], "天台外通向世界真相或规则之外的异常入口。"),
    "dragonfly": ("蜻蜓", ["蜻蜓"], "天台上引导吴翊杰走向透明薄膜的视觉线索。"),
}


DRAMATIC_FUNCTIONS = {
    "1": "建立池清明失踪与他人否认的核心悬念，并引出石琦的旁观线索。",
    "2": "通过网络身份注销和 404 页面强化池清明被系统性抹除。",
    "3": "通过家庭电话否认完成现实关系层面的抹除证据。",
    "4": "以篮球回忆解释“忘掉篮筐”的核心隐喻，并让闫老师替换池清明制造梦魇转场。",
    "5": "在教室和走廊中放大集体遗忘与规则凝视，吴翊杰开始动摇。",
    "6": "石琦揭示记忆错位与消失规则，吴翊杰决定寻找池清明留下的门。",
    "7": "大光以训练和集体惩罚阻拦吴翊杰，规则以日常义务形式出现。",
    "7A": "双马尾女生以受伤和道德压力拖延吴翊杰，制造第二层阻拦。",
    "7B": "楼梯间实体堵塞升级，石琦广播为吴翊杰打开通往天台的缺口。",
    "8": "天台终局对峙，闫老师以升学未来诱惑失败，吴翊杰跃入透明薄膜。",
}


@dataclass
class Paragraph:
    index: int
    text: str


def run(command: list[str], *, capture: bool = True) -> subprocess.CompletedProcess[str]:
    return subprocess.run(command, check=True, text=True, capture_output=capture)


def slug(value: str) -> str:
    return re.sub(r"[^a-zA-Z0-9]+", "_", value).strip("_").lower()


def format_time(seconds: float) -> str:
    seconds = max(0.0, seconds)
    hours = int(seconds // 3600)
    minutes = int((seconds % 3600) // 60)
    whole_seconds = int(seconds % 60)
    millis = int(round((seconds - math.floor(seconds)) * 1000))
    if millis == 1000:
        whole_seconds += 1
        millis = 0
    return f"{hours:02d}:{minutes:02d}:{whole_seconds:02d}.{millis:03d}"


def time_for_filename(seconds: float) -> str:
    return format_time(seconds).replace(":", "-")


def read_paragraphs() -> list[Paragraph]:
    doc = Document(SCRIPT_PATH)
    return [Paragraph(index=i + 1, text=p.text.strip()) for i, p in enumerate(doc.paragraphs) if p.text.strip()]


def classify_segment(text: str) -> dict[str, Any]:
    if text.startswith("出片名") or text.startswith("《") and text.endswith("》"):
        return {"type": "title_card", "text": text}
    if text == "黑屏。" or text == "黑屏":
        return {"type": "blackout", "text": text}
    if text.startswith("【闪回】"):
        return {"type": "flashback", "text": text.replace("【闪回】", "").strip()}
    if text.startswith("【现实】"):
        return {"type": "reality", "text": text.replace("【现实】", "").strip()}
    if "：" in text:
        speaker_part, line = text.split("：", 1)
        speaker = None
        for character in CHARACTER_TERMS.values():
            if any(alias in speaker_part for alias in character["aliases"]):
                speaker = character["name"]
                break
        if not speaker:
            speaker = speaker_part.strip()
        is_os = bool(re.search(r"\bOS\b|os|旁白|女声", speaker_part, re.I))
        return {
            "type": "dialogue_os" if is_os else "dialogue",
            "speaker": speaker,
            "speaker_note": speaker_part.strip(),
            "line": line.strip(),
            "text": text,
        }
    return {"type": "action", "text": text}


def parse_script(paragraphs: list[Paragraph]) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    title = paragraphs[0].text.strip("《》") if paragraphs else "池清明到底是谁"
    scenes: list[dict[str, Any]] = []
    current: dict[str, Any] | None = None

    for paragraph in paragraphs[1:]:
        heading = SCENE_HEADING_RE.match(paragraph.text)
        if heading:
            if current:
                scenes.append(current)
            gd = heading.groupdict()
            current = {
                "id": f"scene_{gd['number']}",
                "scene_number": gd["number"],
                "heading": paragraph.text,
                "location": gd["location"].strip(),
                "time_of_day": gd["time_of_day"],
                "interior_exterior": gd["interior_exterior"],
                "paragraph_start": paragraph.index,
                "paragraph_end": paragraph.index,
                "raw_paragraphs": [],
                "segments": [],
            }
            continue
        if current is None:
            continue
        current["paragraph_end"] = paragraph.index
        segment = classify_segment(paragraph.text)
        segment["paragraph_index"] = paragraph.index
        current["segments"].append(segment)
        current["raw_paragraphs"].append({"paragraph_index": paragraph.index, "text": paragraph.text})

    if current:
        scenes.append(current)

    all_text = "\n".join(p.text for p in paragraphs)
    script = {
        "title": title,
        "logline": "初三学生吴翊杰发现好友池清明从所有人的记忆与记录中消失，在石琦揭示规则后，他选择冲向天台外的异常入口寻找真相。",
        "themes": ["集体遗忘", "规则规训", "青春期反抗", "考试压力", "现实边界与自我选择"],
        "genre": ["校园", "悬疑", "科幻寓言", "青春短片"],
        "source_path": str(SCRIPT_PATH),
        "paragraph_count": len(paragraphs),
        "source_text": all_text,
        "scene_ids": [scene["id"] for scene in scenes],
    }
    return script, scenes


def names_in_text(text: str) -> list[str]:
    found = []
    for character in CHARACTER_TERMS.values():
        if any(alias in text for alias in character["aliases"]):
            found.append(character["name"])
    return sorted(set(found), key=found.index)


def props_in_text(text: str) -> list[str]:
    found = []
    for name, aliases, _description in PROP_TERMS.values():
        if any(alias in text for alias in aliases):
            found.append(name)
    return sorted(set(found), key=found.index)


def build_script_scenes(scenes: list[dict[str, Any]]) -> list[dict[str, Any]]:
    built = []
    for scene in scenes:
        scene_text = "\n".join(item["text"] for item in scene["raw_paragraphs"])
        dialogue_count = sum(1 for segment in scene["segments"] if segment["type"].startswith("dialogue"))
        action_count = sum(1 for segment in scene["segments"] if segment["type"] == "action")
        built.append(
            {
                **scene,
                "dramatic_function": DRAMATIC_FUNCTIONS.get(scene["scene_number"], ""),
                "characters": names_in_text(scene_text),
                "props": props_in_text(scene_text),
                "dialogue_count": dialogue_count,
                "action_count": action_count,
                "summary": summarize_scene(scene),
            }
        )
    return built


def summarize_scene(scene: dict[str, Any]) -> str:
    snippets = [segment.get("line") or segment.get("text", "") for segment in scene["segments"][:4]]
    compact = "；".join(re.sub(r"\s+", "", item) for item in snippets if item)
    return compact[:180]


def evidence_for_terms(paragraphs: list[Paragraph], aliases: list[str], limit: int = 4) -> list[dict[str, Any]]:
    evidence = []
    for paragraph in paragraphs:
        if any(alias in paragraph.text for alias in aliases):
            evidence.append({"paragraph_index": paragraph.index, "matched_text": paragraph.text[:220]})
        if len(evidence) >= limit:
            break
    return evidence


def scene_numbers_for_terms(script_scenes: list[dict[str, Any]], aliases: list[str]) -> list[str]:
    numbers = []
    for scene in script_scenes:
        text = "\n".join(item["text"] for item in scene["raw_paragraphs"])
        if any(alias in text for alias in aliases) or any(alias in scene["location"] for alias in aliases):
            numbers.append(scene["scene_number"])
    return numbers


def build_assets(paragraphs: list[Paragraph], script_scenes: list[dict[str, Any]]) -> list[dict[str, Any]]:
    assets = []
    for key, info in CHARACTER_TERMS.items():
        assets.append(
            {
                "id": f"asset_character_{key}",
                "group": "角色",
                "name": info["name"],
                "aliases": info["aliases"],
                "description": info["description"],
                "appears_in_scenes": scene_numbers_for_terms(script_scenes, info["aliases"]),
                "evidence": evidence_for_terms(paragraphs, info["aliases"]),
                "generation_reference_prompt": info["prompt"],
                "asset_tier": "core" if key in {"wu_yijie", "chi_qingming", "shi_qi", "teacher_yan"} else "supporting",
                "status": "reference_text_ready",
            }
        )
    for key, (name, description) in SCENE_ASSETS.items():
        assets.append(
            {
                "id": f"asset_scene_{key}",
                "group": "场景",
                "name": name,
                "aliases": [name, name.split("·")[-1], name.split("/")[0]],
                "description": description,
                "appears_in_scenes": scene_numbers_for_terms(script_scenes, [name, name.split("·")[-1], name.split("/")[0]]),
                "evidence": evidence_for_terms(paragraphs, [name, name.split("·")[-1], name.split("/")[0]], limit=3),
                "generation_reference_prompt": f"{name}, {description}, Chinese school suspense short film, cinematic lighting",
                "asset_tier": "core",
                "status": "reference_text_ready",
            }
        )
    for key, (name, aliases, description) in PROP_TERMS.items():
        assets.append(
            {
                "id": f"asset_prop_{key}",
                "group": "道具",
                "name": name,
                "aliases": aliases,
                "description": description,
                "appears_in_scenes": scene_numbers_for_terms(script_scenes, aliases),
                "evidence": evidence_for_terms(paragraphs, aliases, limit=3),
                "generation_reference_prompt": f"{name}, {description}, clean prop reference, cinematic realistic style",
                "asset_tier": "core" if key in {"basketball", "borrowing_record", "transparent_membrane"} else "mention",
                "status": "reference_text_ready",
            }
        )
    return assets


def ffprobe_metadata() -> dict[str, Any]:
    result = run(
        [
            "ffprobe",
            "-v",
            "error",
            "-show_entries",
            "format=duration,size,bit_rate:stream=index,codec_type,codec_name,width,height,avg_frame_rate,duration",
            "-of",
            "json",
            str(VIDEO_PATH),
        ]
    )
    return json.loads(result.stdout)


def detect_scene_cuts() -> list[float]:
    command = [
        "ffmpeg",
        "-hide_banner",
        "-nostats",
        "-i",
        str(VIDEO_PATH),
        "-vf",
        f"select='gt(scene,{SCENE_THRESHOLD})',showinfo",
        "-an",
        "-f",
        "null",
        "-",
    ]
    proc = subprocess.run(command, check=True, text=True, capture_output=True)
    times = []
    for line in proc.stderr.splitlines():
        match = re.search(r"pts_time:([0-9.]+)", line)
        if match:
            value = float(match.group(1))
            if value >= MIN_SHOT_SECONDS:
                times.append(value)
    deduped = []
    for value in sorted(times):
        if not deduped or value - deduped[-1] >= MIN_SHOT_SECONDS:
            deduped.append(value)
    return deduped


def build_boundaries(cuts: list[float], duration: float) -> list[float]:
    boundaries = [0.0] + [cut for cut in cuts if MIN_SHOT_SECONDS <= cut < duration - MIN_SHOT_SECONDS] + [duration]
    boundaries = sorted(set(round(item, 3) for item in boundaries))
    merged = [boundaries[0]]
    for boundary in boundaries[1:]:
        if boundary - merged[-1] < MIN_SHOT_SECONDS and boundary != duration:
            continue
        merged.append(boundary)
    if merged[-1] != duration:
        merged.append(duration)
    return merged


def weighted_scene_ranges(script_scenes: list[dict[str, Any]], duration: float) -> list[dict[str, Any]]:
    weights = [max(1, len(scene["segments"])) for scene in script_scenes]
    total = sum(weights)
    cursor = 0.0
    ranges = []
    for index, scene in enumerate(script_scenes):
        segment_duration = duration * weights[index] / total
        end = duration if index == len(script_scenes) - 1 else cursor + segment_duration
        ranges.append({"scene": scene, "start": cursor, "end": end})
        cursor = end
    return ranges


def scene_for_time(scene_ranges: list[dict[str, Any]], midpoint: float) -> dict[str, Any]:
    for item in scene_ranges:
        if item["start"] <= midpoint < item["end"]:
            return item["scene"]
    return scene_ranges[-1]["scene"]


def export_frame(time_seconds: float, output_path: Path) -> None:
    if output_path.exists():
        return
    output_path.parent.mkdir(parents=True, exist_ok=True)
    run(
        [
            "ffmpeg",
            "-hide_banner",
            "-loglevel",
            "error",
            "-ss",
            f"{time_seconds:.3f}",
            "-i",
            str(VIDEO_PATH),
            "-frames:v",
            "1",
            "-q:v",
            "3",
            "-vf",
            "scale=960:-2",
            str(output_path),
        ],
        capture=True,
    )


def build_edit_shots(boundaries: list[float], script_scenes: list[dict[str, Any]]) -> list[dict[str, Any]]:
    duration = boundaries[-1]
    scene_ranges = weighted_scene_ranges(script_scenes, duration)
    shots = []
    for index, (start, end) in enumerate(zip(boundaries, boundaries[1:]), start=1):
        shot_duration = round(end - start, 3)
        midpoint = start + shot_duration / 2
        scene = scene_for_time(scene_ranges, midpoint)
        keyframe_time = min(max(start + min(shot_duration / 2, 1.5), start + 0.05), max(start, end - 0.05))
        keyframe_name = f"shot_{index:04d}_{time_for_filename(start)}.jpg"
        keyframe_path = KEYFRAME_DIR / keyframe_name
        export_frame(keyframe_time, keyframe_path)

        supplemental = []
        alignment_status = "auto_aligned_needs_review"
        quality_notes = []
        if shot_duration > LONG_SHOT_REVIEW_SECONDS:
            alignment_status = "needs_manual_review"
            quality_notes.append(f"Shot duration {shot_duration:.2f}s exceeds {LONG_SHOT_REVIEW_SECONDS:.0f}s; likely contains internal action or missed cuts.")
            for marker, ratio in [("mid", 0.5), ("late", 0.8)]:
                supplemental_time = start + shot_duration * ratio
                supplemental_name = f"shot_{index:04d}_{marker}_{time_for_filename(supplemental_time)}.jpg"
                supplemental_path = KEYFRAME_DIR / supplemental_name
                export_frame(supplemental_time, supplemental_path)
                supplemental.append({"time": round(supplemental_time, 3), "path": str(supplemental_path)})

        scene_text = "\n".join(item["text"] for item in scene["raw_paragraphs"][:5])
        shot = {
            "id": f"edit_shot_{index:04d}",
            "shot_number": index,
            "start": round(start, 3),
            "end": round(end, 3),
            "start_timecode": format_time(start),
            "end_timecode": format_time(end),
            "duration": shot_duration,
            "script_scene_id": scene["id"],
            "script_scene_number": scene["scene_number"],
            "script_scene_heading": scene["heading"],
            "script_evidence": scene_text[:360],
            "visual_summary": f"自动切点镜头；粗对齐到剧本场景 {scene['scene_number']}「{scene['location']}」。画面内容需结合关键帧人工复核。",
            "dialogue_or_audio": dialogue_summary(scene),
            "camera": "unknown_needs_visual_review",
            "transition": "first_frame" if index == 1 else "detected_cut",
            "characters": scene["characters"],
            "props": scene["props"],
            "scene": scene["location"],
            "keyframe_path": str(keyframe_path),
            "keyframe_time": round(keyframe_time, 3),
            "supplemental_keyframes": supplemental,
            "alignment_status": alignment_status,
            "confidence": "medium" if alignment_status == "auto_aligned_needs_review" else "low",
            "quality_notes": quality_notes,
        }
        shots.append(shot)
    return shots


def dialogue_summary(scene: dict[str, Any]) -> str:
    lines = []
    for segment in scene["segments"]:
        if segment["type"].startswith("dialogue"):
            speaker = segment.get("speaker", "")
            line = segment.get("line", "")
            lines.append(f"{speaker}: {line}")
        if len(lines) >= 3:
            break
    return " / ".join(lines) if lines else "无明确对白；以动作、环境声或音乐为主。"


def build_filmforge_mapping(script: dict[str, Any], assets: list[dict[str, Any]], edit_shots: list[dict[str, Any]]) -> dict[str, Any]:
    mapped_assets = [
        {
            "id": asset["id"],
            "group": asset["group"],
            "name": asset["name"],
            "description": asset["description"],
            "image": "",
            "prompt": asset["generation_reference_prompt"],
            "aliases": asset.get("aliases", []),
            "evidence": asset.get("evidence", []),
            "assetTier": asset.get("asset_tier"),
            "status": "queued",
        }
        for asset in assets
    ]
    mapped_shots = [
        {
            "id": shot["shot_number"],
            "unitType": "key_image",
            "title": f"镜头 {shot['shot_number']:04d}",
            "duration": f"{shot['duration']:.1f}s",
            "durationSeconds": shot["duration"],
            "description": shot["visual_summary"],
            "sceneNumber": shot["script_scene_number"],
            "sceneHeading": shot["script_scene_heading"],
            "sourceText": shot["script_evidence"],
            "interiorExterior": "UNKNOWN",
            "timeOfDay": "",
            "scene": shot["scene"],
            "characters": shot["characters"],
            "props": shot["props"],
            "image": shot["keyframe_path"],
            "imageStatus": "succeeded",
            "videoStatus": "succeeded",
            "videoUrl": str(VIDEO_PATH),
            "videoAttempts": 1,
            "reviewStatus": "pending",
        }
        for shot in edit_shots
    ]
    timeline = [
        {
            "id": f"timeline_{shot['shot_number']:04d}",
            "shotId": shot["shot_number"],
            "label": f"镜头 {shot['shot_number']:04d}",
            "start": shot["start"],
            "end": shot["end"],
            "duration": shot["duration"],
        }
        for shot in edit_shots
    ]
    return {
        "project": {
            "id": "project_chi_qingming_golden_case",
            "title": script["title"],
            "stage": "review_ready",
            "brief": {
                "title": script["title"],
                "logline": script["logline"],
                "audience": "AIGC 影视工作台内部样本、剧本到视频评测集",
                "tone": "校园悬疑、现实规训、科幻寓言",
                "format": "AI 短片",
                "durationTarget": "约 14 分钟",
            },
            "settings": {
                "locale": "zh-CN",
                "aspectRatio": "16:9",
                "resolution": "720p",
                "imageModel": "dataset-reference",
                "assetModel": "dataset-reference",
                "videoModel": "final-cut-reference",
                "workflow": "first_frame",
                "visualStyle": "写实电影感",
                "audioMode": "auto",
            },
        },
        "assets": mapped_assets,
        "shots": mapped_shots,
        "timeline": timeline,
    }


def validate_dataset(dataset: dict[str, Any]) -> list[str]:
    errors = []
    scene_numbers = {scene["scene_number"] for scene in dataset["script_scenes"]}
    expected = {"1", "2", "3", "4", "5", "6", "7", "7A", "7B", "8"}
    missing = expected - scene_numbers
    if missing:
        errors.append(f"Missing script scenes: {sorted(missing)}")
    shots = dataset["edit_shots"]
    if not shots:
        errors.append("No edit shots generated.")
    else:
        if abs(shots[0]["start"]) > 0.01:
            errors.append("First shot does not start at 0.")
        video_duration = float(dataset["metadata"]["video"]["format"]["duration"])
        if abs(shots[-1]["end"] - video_duration) > 0.1:
            errors.append("Last shot does not end near video duration.")
        for left, right in zip(shots, shots[1:]):
            if abs(left["end"] - right["start"]) > 0.02:
                errors.append(f"Timeline gap/overlap between {left['id']} and {right['id']}.")
                break
    asset_names = {asset["name"] for asset in dataset["assets"]}
    for shot in shots:
        for name in shot["characters"] + shot["props"]:
            if name not in asset_names:
                errors.append(f"{shot['id']} references missing asset {name}.")
    return errors


def main() -> None:
    ROOT.mkdir(parents=True, exist_ok=True)
    KEYFRAME_DIR.mkdir(parents=True, exist_ok=True)
    paragraphs = read_paragraphs()
    script, raw_scenes = parse_script(paragraphs)
    script_scenes = build_script_scenes(raw_scenes)
    assets = build_assets(paragraphs, script_scenes)
    video = ffprobe_metadata()
    video_duration = float(video["format"]["duration"])
    cuts = detect_scene_cuts()
    boundaries = build_boundaries(cuts, video_duration)
    edit_shots = build_edit_shots(boundaries, script_scenes)
    dataset = {
        "metadata": {
            "dataset_id": "chi_qingming_who_is_chi_qingming",
            "project_name": "FilmForge 黄金样本：池清明到底是谁",
            "title": script["title"],
            "created_at": datetime.now().isoformat(timespec="seconds"),
            "schema_version": "0.1.0",
            "source_files": {
                "script_docx": str(SCRIPT_PATH),
                "final_cut_video": str(VIDEO_PATH),
            },
            "video": video,
            "processing": {
                "scene_detection": {
                    "tool": "ffmpeg select gt(scene,threshold)",
                    "threshold": SCENE_THRESHOLD,
                    "raw_cut_count": len(cuts),
                    "merged_shot_count": len(edit_shots),
                    "min_shot_seconds": MIN_SHOT_SECONDS,
                    "long_shot_review_seconds": LONG_SHOT_REVIEW_SECONDS,
                },
                "notes": [
                    "Shot boundaries are automatic candidates and need human review for final edit-grade labeling.",
                    "Dialogue/audio fields are derived from script text, not speech recognition.",
                    "Visual summaries are coarse script-aligned placeholders until keyframes are manually reviewed.",
                ],
            },
        },
        "script": script,
        "assets": assets,
        "script_scenes": script_scenes,
        "edit_shots": edit_shots,
        "filmforge_mapping": build_filmforge_mapping(script, assets, edit_shots),
        "quality_notes": [],
    }
    validation_errors = validate_dataset(dataset)
    dataset["quality_notes"] = [
        {
            "type": "validation",
            "status": "passed" if not validation_errors else "needs_attention",
            "messages": validation_errors,
        },
        {
            "type": "manual_review",
            "status": "required",
            "messages": [
                "Review keyframes to replace visual_summary/camera with observed content.",
                "Confirm script_scene_id for every edit_shot before using this as a training label.",
            ],
        },
    ]
    DATASET_PATH.write_text(json.dumps(dataset, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({
        "dataset": str(DATASET_PATH),
        "script_scene_count": len(script_scenes),
        "asset_count": len(assets),
        "raw_cut_count": len(cuts),
        "edit_shot_count": len(edit_shots),
        "validation_errors": validation_errors,
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
